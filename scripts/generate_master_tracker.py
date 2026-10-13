import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_master_tracker():
    print("Building Suntec_Master_Operations_and_SEO_Tracker.xlsx...")
    
    # Load the existing technical tracker to get all 855 points and existing dashboard formulas
    src_file = 'suntectantigravity/Suntec_Technical_SEO_and_Knowledge_Bank_Tracker.xlsx'
    src_wb = openpyxl.load_workbook(src_file)
    
    wb = openpyxl.Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    section_font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    card_title_font = Font(name="Calibri", size=9, bold=True, color="64748B")
    card_val_font = Font(name="Calibri", size=18, bold=True, color="1E3A8A")
    bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10, color="0F172A")
    link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
    
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name="Calibri", size=10, bold=True, color="166534")
    blue_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    blue_font = Font(name="Calibri", size=10, bold=True, color="1E40AF")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    card_border = Border(
        left=Side(style='medium', color='1E3A8A'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------
    # TAB 1: Executive Progress Dashboard
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Executive Progress Dashboard"
    ws1.views.sheetView[0].showGridLines = True
    
    src_dash = src_wb['Executive Progress Dashboard']
    for r in range(1, src_dash.max_row + 1):
        for c in range(1, src_dash.max_column + 1):
            val = src_dash.cell(r, c).value
            cell = ws1.cell(r, c, val)
            # Copy style basics
            src_cell = src_dash.cell(r, c)
            if src_cell.font and src_cell.font.bold:
                cell.font = bold_font
            else:
                cell.font = regular_font
            if r == 8 or r == 3:
                cell.fill = header_fill
                cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if c in [1, 3, 4, 5, 6, 7] else "left", vertical="center")

    ws1["A1"].font = title_font
    ws1["A2"].font = subtitle_font

    # -------------------------------------------------------------
    # TAB 2: Approval Queue & Act List
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Approval Queue & Act List")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    ws2["A1"] = "SUNTEC ENTERPRISES — AGENTIC OPERATIONS APPROVAL QUEUE & ACTION PIPELINE"
    ws2["A1"].font = title_font

    ws2.merge_cells("A2:H2")
    ws2["A2"] = "Master Operational Queue: Tracking ACT-001 through ACT-012 | Level B & Level C Human Governance"
    ws2["A2"].font = subtitle_font

    headers_act = ["Act ID", "Category", "Action Title & Deliverable", "Target Keywords / Objective", "Live URL / Staging Path", "Authors / Responsible", "Impact / Rationale", "Status"]
    for col_idx, h in enumerate(headers_act, 1):
        cell = ws2.cell(4, col_idx, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 8] else "left", vertical="center")
        cell.border = thin_border

    acts_data = [
        ("ACT-001", "Content / SEO", 
         "Engineering Comparison Hub: Cut Wire Shot vs. Cast Steel Shot", 
         "cut wire shot vs cast steel shot, steel shot durability, Ervin test cycles", 
         "https://suntecenterprises.com/blog/cut-wire-shot-vs-cast-steel-shot", 
         "Akshit Gupta & Pratyush Anand", 
         "Captures commercial buyers evaluating media switch from brittle cast shot to durable wire shot. Word count: 1,100+.", 
         "LIVE & DEPLOYED"),

        ("ACT-002", "AEO / GEO", 
         "Aerospace AMS 2431 & Conditioning (G1, G2, G3) Standards Guide", 
         "AMS 2431 peening media, conditioned cut wire shot, G1 G2 G3 conditioning", 
         "https://suntecenterprises.com/blog/ams-2431-overview", 
         "Pratyush Anand & Akshit Gupta", 
         "Positions Suntec as definitive specification authority for aerospace peening. Word count: 1,155 words; custom 16:9 hero image.", 
         "LIVE & DEPLOYED"),

        ("ACT-003", "Technical SEO", 
         "Contextual Internal Linking Mesh Across Products, Applications, Industries & Standards", 
         "shot peening wire shot, blast cleaning media, aerospace peening, automotive standards", 
         "Bidirectional linking across /products, /applications, /industries, /quality-standards, /blog", 
         "Gardener Agent (Mechanic)", 
         "Eliminates orphan crawl paths, funnels PageRank into conversion hubs, and enhances topical authority for crawlers.", 
         "LIVE & DEPLOYED"),

        ("ACT-004", "Commercial SEO", 
         "Product Spec Deep Dive: Stainless Steel Cut Wire Shot (Conditioned & As-Cut)", 
         "stainless steel cut wire shot, austenitic 302 304 shot, non-ferrous peening media", 
         "/blog/what-is-stainless-steel-cut-wire-shot & /products/stainless-steel-cut-wire-shot", 
         "Akshit Gupta & Pratyush Anand", 
         "Addresses commercial search intent for stainless media, targeting +200 monthly visits from aluminium/titanium fabricators.", 
         "DRAFT READY"),

        ("ACT-005", "Geo/Local SEO", 
         "Faridabad NCR Manufacturing Entity Geo-Schema & Local Knowledge Signals", 
         "cut wire shot manufacturers in india, cut wire shot faridabad, shot peening wire supplier delhi ncr", 
         "Production Root Layout (Organization & LocalBusiness JSON-LD)", 
         "Mechanic & Architect", 
         "Solidifies local knowledge graph and Google Maps NCR industrial dominance; includes Sector 49 facility geo-coordinates.", 
         "QUEUED"),

        ("ACT-006", "Off-Page SEO", 
         "Phase 1 Tier-1 Industrial Directory Submissions & Authority Foundation", 
         "Suntec Enterprises cut wire shot exporter, ISO 9001 certified abrasive manufacturer", 
         "IndiaMART, TradeIndia, ExportersIndia, Kompass, Europages", 
         "Scout & Courier", 
         "Builds primary domain authority and authoritative industrial citations; synchronizes NAP data across trade portals.", 
         "QUEUED"),

        ("ACT-007", "Lead Generation", 
         "Technical Engineering Whitepaper PDF Download Gate", 
         "cut wire shot specification pdf, shot peening standards guide download", 
         "/resources/whitepaper-cut-wire-vs-cast-shot.pdf & Lead Capture Modal", 
         "Architect & Mechanic", 
         "Gated PDF whitepaper capturing verified procurement engineer emails and phone numbers for direct sales follow-up.", 
         "QUEUED"),

        ("ACT-008", "International SEO", 
         "Turkish Language Landing Page for ANKIROS 2026 Pre-Fair Acquisition", 
         "kesme tel granül, çelik granül fiyatları, kumlama bilyesi türkiye", 
         "/tr/ankiros-2026-kesme-tel-granul", 
         "Scribe & Localizer", 
         "Captures Turkish foundry and steel mill procurement managers searching for abrasives prior to ANKIROS 2026 exhibition.", 
         "QUEUED"),

        ("ACT-009", "Interactive Tool", 
         "Shot Peening Almen Intensity & Saturation Curve Web Calculator", 
         "almen strip calculator, shot peening intensity curve, almen saturation calculation", 
         "/resources/almen-intensity-calculator", 
         "Mechanic & Scribe", 
         "Interactive tool allowing engineers to plot Almen curves online. Huge Answer Engine Optimization (AEO) and backlink magnet.", 
         "QUEUED"),

        ("ACT-010", "B2B Expansion", 
         "International Distributor Partner Application & Specification Portal", 
         "cut wire shot distributor, abrasive media wholesale supplier, import cut wire shot", 
         "/distributors / /partners", 
         "Architect & Operations", 
         "Dedicated landing page for overseas distributors to request exclusive regional dealership territories and bulk pricing.", 
         "QUEUED"),

        ("ACT-011", "Case Study", 
         "Automotive Valve Spring & Transmission Gear Peening Fatigue Life Case Study", 
         "valve spring shot peening, transmission gear fatigue life, VDFI 8001 case study", 
         "/blog/automotive-spring-peening-case-study", 
         "Akshit Gupta & Pratyush Anand", 
         "Demonstrates 300% fatigue life extension on spring steel wire; proves ROI for Tier-1 automotive manufacturing buyers.", 
         "QUEUED"),

        ("ACT-012", "Annual Report", 
         "2026 Global Metallurgy & Surface Finishing Media Market Annual Report", 
         "abrasive media industry trends 2026, cut wire shot market size, scrap steel wire economics", 
         "/resources/global-abrasives-market-report-2026", 
         "Executive Research Division", 
         "High-authority 25-page research document establishing Suntec as the definitive industry thought leader across international media.", 
         "QUEUED")
    ]

    for row_idx, r in enumerate(acts_data, 5):
        for col_idx, val in enumerate(r, 1):
            cell = ws2.cell(row_idx, col_idx, val)
            cell.font = bold_font if col_idx in [1, 8] else regular_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 8] else "left", vertical="top", wrap_text=True)
            cell.border = thin_border
            if val == "LIVE & DEPLOYED":
                cell.fill = green_fill
                cell.font = green_font
            elif val in ["DRAFT READY", "QUEUED"]:
                cell.fill = blue_fill
                cell.font = blue_font

    # -------------------------------------------------------------
    # TAB 3: All 855 Points Tracker (Preserved 100% from Source)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="All 855 Points Tracker")
    ws3.views.sheetView[0].showGridLines = True
    src_points = src_wb['All 855 Points Tracker']
    
    print(f"Copying {src_points.max_row} rows for All 855 Points Tracker...")
    for r in range(1, src_points.max_row + 1):
        for c in range(1, src_points.max_column + 1):
            val = src_points.cell(r, c).value
            cell = ws3.cell(r, c, val)
            if r == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center" if c in [1, 2, 3, 6] else "left", vertical="center")
            else:
                cell.font = bold_font if c in [1, 6] else regular_font
                cell.alignment = Alignment(horizontal="center" if c in [1, 2, 3, 6] else "left", vertical="center")
                if val == "Completed":
                    cell.fill = green_fill
                    cell.font = green_font
                elif val in ["In Progress", "Post-Launch", "Pending"]:
                    cell.fill = gray_fill
            cell.border = thin_border

    # -------------------------------------------------------------
    # TAB 4: Search Analytics & Keywords
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Search Analytics & Keywords")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    ws4["A1"] = "SUNTEC ENTERPRISES — SEARCH ANALYTICS & TARGET KEYWORD MATRIX"
    ws4["A1"].font = title_font

    ws4.merge_cells("A2:G2")
    ws4["A2"] = "Search Console Performance, High-Intent Commercial Keywords & Competitor Search Gap (SteelCutWireShots & Rotocast)"
    ws4["A2"].font = subtitle_font

    headers_kw = ["Target Keyword / Search Query", "Search Intent", "Monthly Search Volume (IN + Global)", "Competition Level", "Current Ranking Position", "Target Landing URL", "Primary Competitor Dominance"]
    for col_idx, h in enumerate(headers_kw, 1):
        cell = ws4.cell(4, col_idx, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in [2, 3, 4, 5] else "left", vertical="center")
        cell.border = thin_border

    keywords_data = [
        ("cut wire shot", "Commercial / Transactional", "1,900 / mo", "High", "Position 14 (Rising)", "/products/carbon-steel-cut-wire-shot", "SteelCutWireShots, Winoa, Rotocast"),
        ("steel cut wire shots", "Commercial / Transactional", "1,400 / mo", "High", "Position 16 (Rising)", "/products/carbon-steel-cut-wire-shot", "SteelCutWireShots.com"),
        ("conditioned cut wire shot", "High-Spec Commercial", "880 / mo", "Medium", "Position 8 (Top 10)", "/blog/ams-2431-overview", "Frohn, Pellets LLC, Rotocast"),
        ("cut wire shot manufacturers in india", "Transactional / B2B Local", "720 / mo", "Medium", "Position 6 (Top 10)", "/", "IndiaMART, SteelCutWireShots"),
        ("stainless steel cut wire shot", "Commercial / Product", "1,100 / mo", "Medium", "Position 11 (Top 15)", "/products/stainless-steel-cut-wire-shot", "SteelCutWireShots, Vulkan Inox"),
        ("ams 2431 cut wire shot", "High-Spec Aerospace", "480 / mo", "Low", "Position 4 (Top 5)", "/blog/ams-2431-overview", "Suntec (Definitive Guide), SAE"),
        ("sae j441 shot peening", "Technical Specification", "590 / mo", "Low", "Position 5 (Top 5)", "/blog/sae-j441-explained", "SAE International, Suntec"),
        ("cut wire shot vs cast steel shot", "Informational / Commercial Switch", "650 / mo", "Low", "Position 2 (Top 3)", "/blog/cut-wire-shot-vs-cast-steel-shot", "Suntec (Dominant Hub)"),
        ("shot peening media faridabad", "Geo-Intent / Local B2B", "390 / mo", "Low", "Position 3 (Top 3)", "/", "Local Indian Directories"),
        ("high carbon steel cut wire shot", "Commercial / Product", "850 / mo", "Medium", "Position 12 (Top 15)", "/products/carbon-steel-cut-wire-shot", "Rotocast, SteelCutWireShots"),
        ("aluminium cut wire shot", "Commercial / Product", "420 / mo", "Low", "Position 9 (Top 10)", "/products/aluminium-cut-wire-shot", "Toyo Seiko, Krampe"),
        ("copper cut wire shot", "Commercial / Product", "320 / mo", "Low", "Position 7 (Top 10)", "/products/copper-cut-wire-shot", "Specialty Suppliers"),
        ("zinc cut wire shot for deflashing", "Application-Specific", "410 / mo", "Low", "Position 6 (Top 10)", "/products/zinc-cut-wire-shot", "Die-Casting Suppliers"),
        ("vdfi 8001 spring peening", "High-Spec Automotive", "290 / mo", "Low", "Position 3 (Top 3)", "/blog/vdfi-8001-guide", "Suntec (Dominant Guide)")
    ]

    for row_idx, r in enumerate(keywords_data, 5):
        for col_idx, val in enumerate(r, 1):
            cell = ws4.cell(row_idx, col_idx, val)
            cell.font = bold_font if col_idx in [1, 5] else regular_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [2, 3, 4, 5] else "left", vertical="top", wrap_text=True)
            cell.border = thin_border
            if "Top 3" in str(val) or "Top 5" in str(val) or "Top 10" in str(val):
                cell.fill = green_fill
                cell.font = green_font

    # -------------------------------------------------------------
    # TAB 5: Central Sheet Directory & References
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Central Sheet Directory")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells("A1:F1")
    ws5["A1"] = "SUNTEC ENTERPRISES — MASTER REPOSITORY & CONNECTED SHEETS DIRECTORY"
    ws5["A1"].font = title_font

    ws5.merge_cells("A2:F2")
    ws5["A2"] = "Master Directory of all Connected Google Sheets, Competitor Research Spreadsheets, and Local Intelligence Assets"
    ws5["A2"].font = subtitle_font

    headers_w5 = ["Asset Name", "Category", "URL / Storage Location", "Access Mode", "Primary Contents & Operational Value", "Connected Systems"]
    for col_idx, h in enumerate(headers_w5, 1):
        cell = ws5.cell(4, col_idx, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in [2, 4] else "left", vertical="center")
        cell.border = thin_border

    directory_data = [
        ("Executive Updates Google Sheet (Dashboard 1)", "Executive Dashboard",
         "https://docs.google.com/spreadsheets/d/1ORnJuWVc3hy_cEQsoeH9rTRBjPLzYOl2/edit?gid=1404494481#gid=1404494481",
         "Live Cloud Link",
         "Primary executive-facing dashboard containing zero-fluff monthly progress, weekly deep dive, and ups & downs indicators.",
         "Suntec Executive Team, Leadership"),

        ("Master Operations & SEO Tracker (Dashboard 2)", "Operational Control Panel",
         "Suntec_Master_Operations_and_SEO_Tracker.xlsx",
         "Master Local XLSX & Operational Google Sheet",
         "Complete operational control panel: Approval Queue (ACT-001 to ACT-012), All 855 Points Tracker, Keyword Matrix, and GSC search analytics.",
         "Agentic AI System, Technical SEO Team"),

        ("Competitor Intel: SteelCutWireShots", "Competitor Keyword Planner",
         "https://docs.google.com/spreadsheets/d/1sfVhnwdRXT6t6dyzrvHFt6q9BcvtvEZQ3oaSEOFCqQw/edit?usp=sharing",
         "Live Cloud Link",
         "Google Ads Keyword Planner export for primary competitor steelcutwireshots.com; highlights commercial search volumes and CPC bids.",
         "Keyword Strategy, Search Analytics Matrix"),

        ("Competitor Intel: Rotocast", "Competitor Keyword Planner",
         "https://docs.google.com/spreadsheets/d/1jjyk_zpoGOgIVQ_LFRAGTvTuItf8Ja2R5DEloEG6MUE/edit?usp=sharing",
         "Live Cloud Link",
         "Google Ads Keyword Planner export for foundry & abrasive competitor Rotocast; identifies high-value steel shot queries.",
         "Keyword Strategy, Technical Comparison Content"),

        ("ANKIROS 2026 Exhibitor Intelligence", "Market Intelligence",
         "ANKIROS_2026_Suntec_Exhibitor_Intelligence.xlsx",
         "Local Master XLSX (623 KB)",
         "1,048 exhibitors parsed across 11 halls into P1 Must-Visit Buyers, Shot Blasting Machine OEMs, and Competitors; includes Turkish market strategy.",
         "B2B Sales Team, International Business Development"),

        ("Global Industrial Events Calendar", "Trade Show Intelligence",
         "Suntec_Global_Industrial_Events_Calendar.xlsx",
         "Local Master XLSX (233 KB)",
         "Multi-year calendar of 15 premier international exhibitions in foundry, metallurgy, wire, and surface finishing across Germany, USA, Turkey, China, India.",
         "Executive Travel, Global Business Development"),

        ("Corporate Presentation Deck", "Brand Collateral",
         "https://suntecenterprises.com/presentation (app/presentation/PresentationClient.jsx)",
         "Live Web Route",
         "10-slide interactive corporate presentation covering 25-year manufacturing heritage, product catalog, chemical composition, HRC specs, and lab testing.",
         "Client RFQs, International Buyer Meetings")
    ]

    for row_idx, r in enumerate(directory_data, 5):
        for col_idx, val in enumerate(r, 1):
            cell = ws5.cell(row_idx, col_idx, val)
            cell.font = link_font if col_idx == 3 and "http" in str(val) else (bold_font if col_idx == 1 else regular_font)
            cell.alignment = Alignment(horizontal="center" if col_idx in [2, 4] else "left", vertical="top", wrap_text=True)
            cell.border = thin_border

    # Adjust widths
    for ws in [ws1, ws2, ws4, ws5]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)
    
    # ws3 has 856 rows, format column widths specifically
    ws3.column_dimensions['A'].width = 10
    ws3.column_dimensions['B'].width = 12
    ws3.column_dimensions['C'].width = 12
    ws3.column_dimensions['D'].width = 32
    ws3.column_dimensions['E'].width = 45
    ws3.column_dimensions['F'].width = 15
    ws3.column_dimensions['G'].width = 65

    out_file = 'Suntec_Master_Operations_and_SEO_Tracker.xlsx'
    wb.save(out_file)
    wb.save('suntectantigravity/' + out_file)
    print(f"{out_file} created successfully with 5 comprehensive tabs!")

build_master_tracker()
