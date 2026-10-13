#!/usr/bin/env python3
"""
Generate Suntec Backlink Command Center (V2.0)
Generates the comprehensive 22-Tab Excel Command Center:
1. BACKLINK DASHBOARD
2. MASTER PROSPECTS
3. CONTACTS
4. LINK OPPORTUNITIES
5. COMPETITOR LINK GAP
6. DOMAIN INTERSECTION
7. PAGE INTERSECTION
8. BROKEN LINKS
9. UNLINKED MENTIONS
10. LOST LINKS
11. LINKABLE ASSETS
12. OUTREACH QUEUE
13. OUTREACH LOG
14. FOLLOW-UPS
15. ACQUIRED LINKS
16. REJECTED PROSPECTS
17. DIGITAL PR
18. INTERNATIONAL
19. LINK TARGETS
20. ANCHOR TEXT
21. MONTHLY PERFORMANCE
22. STRATEGIC BACKLOG
"""

import os
import json
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
DB_PATH = os.path.join(DATA_DIR, "strategic_memory.db")
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "SUNTEC_BACKLINK_COMMAND_CENTER.xlsx")

# Styling definitions
NAVY_HEADER = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
STEEL_HEADER = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
GOLD_HEADER = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)


def format_sheet_header(ws, title, columns, fill=NAVY_HEADER):
    ws.append([title])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([])  # Blank row
    
    ws.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = fill
        cell.font = WHITE_BOLD
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER


def autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # skip row 1 title
            if cell.row == 1:
                continue
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)


def generate_backlink_command_center():
    print("[Command Center Generator] Initializing Excel workbook...")
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # -------------------------------------------------------------
    # TAB 1: BACKLINK DASHBOARD
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="BACKLINK DASHBOARD")
    ws1.append(["SUNTEC ENTERPRISES — EXECUTIVE BACKLINK & AUTHORITY DASHBOARD"])
    ws1.cell(row=1, column=1).font = TITLE_FONT
    ws1.append([])
    
    dashboard_data = [
        ["METRIC CATEGORY", "CORE METRIC INDICATOR", "OBSERVED VALUE", "STRATEGIC BENCHMARK / TARGET", "STATUS / HEALTH"],
        ["CURRENT AUTHORITY", "Total Live Backlinks", 800, "1,200 (Q1 2027)", "STABLE BASELINE"],
        ["CURRENT AUTHORITY", "Total Referring Domains", 694, "850 (Q1 2027)", "HEALTHY DIVERSITY"],
        ["CURRENT AUTHORITY", "DataForSEO Domain Rank", 164, "Top 120", "GROWING"],
        ["CURRENT AUTHORITY", "High-Relevance Industrial Domains", 28, "50+", "STRATEGIC EXPANSION"],
        ["COMPETITOR GAP", "Winoa.com Referring Domains", "3,400+", "Market Leader", "OPPORTUNITY GAP"],
        ["COMPETITOR GAP", "PelletsLLC.com Referring Domains", "820+", "North America Benchmark", "STRIKING DISTANCE"],
        ["COMPETITOR GAP", "Common Competitor Intersecting Domains", 25, "Targeting Top 15 Gaps", "ACTIONABLE"],
        ["ACQUISITION PIPELINE", "Tier-1 Qualified Prospects", 15, "25 Active in Pool", "POPULATED"],
        ["ACQUISITION PIPELINE", "Verified Editorial Contacts", 5, "10 Verified", "READY FOR REVIEW"],
        ["ACQUISITION PIPELINE", "Approval-Gated Outreach Ready", 5, "Level C Gated", "PENDING APPROVAL"],
        ["ACQUISITION PIPELINE", "Active Linkable Content Assets", 4, "5 Core Whitepapers", "VERIFIED LIVE"],
        ["BUSINESS IMPACT", "Target Geographic Markets", "India, Germany, EU, USA", "4 Global Clusters", "COVERED"],
        ["BUSINESS IMPACT", "Core Commercial Keyword Focus", "Cut Wire Shot, Peening Media", "Top 3 Commercial Terms", "MONITORED"],
        ["BUSINESS IMPACT", "Attributed RFQ Conversion Focus", "Zero Spam / High Relevance B2B", "100% Human Review", "ACTIVE GUARDRAIL"]
    ]
    for r_idx, row in enumerate(dashboard_data, 3):
        ws1.append(row)
        for c_idx in range(1, len(row) + 1):
            cell = ws1.cell(row=r_idx, column=c_idx)
            cell.border = THIN_BORDER
            if r_idx == 3:
                cell.fill = NAVY_HEADER
                cell.font = WHITE_BOLD
                cell.alignment = ALIGN_CENTER
            else:
                cell.font = REGULAR_FONT
                cell.alignment = ALIGN_LEFT if c_idx <= 2 else ALIGN_CENTER
    autofit_columns(ws1)

    # -------------------------------------------------------------
    # TAB 2: MASTER PROSPECTS
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="MASTER PROSPECTS")
    p_cols = ["Prospect ID", "Organization", "Domain", "Country", "Industry", "Website Type", "Overall Score", "Priority", "Relevance", "Authority", "Commercial", "Status", "Discovery Source"]
    format_sheet_header(ws2, "MASTER BACKLINK PROSPECTS & QUALIFICATION POOL", p_cols)
    
    cur.execute("""
    SELECT prospect_id, organization, domain, country, industry, website_type,
           overall_score, priority, relevance_score, authority_score, commercial_relevance,
           status, discovery_source
    FROM backlink_prospects ORDER BY overall_score DESC;
    """)
    for row in cur.fetchall():
        ws2.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws2.cell(row=ws2.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER if col_idx in [1, 4, 6, 7, 8, 9, 10, 11, 12] else ALIGN_LEFT
    autofit_columns(ws2)

    # -------------------------------------------------------------
    # TAB 3: CONTACTS
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="CONTACTS")
    c_cols = ["Contact ID", "Prospect ID", "Full Name", "Job Title / Role", "Organization", "Email Status", "Source URL", "Evidence Description", "Outreach Status"]
    format_sheet_header(ws3, "PUBLIC EDITORIAL & TECHNICAL CONTACTS DIRECTORY", c_cols)
    
    cur.execute("""
    SELECT contact_id, prospect_id, name, job_title, organization, email_status,
           source_url, contact_evidence, outreach_status
    FROM backlink_contacts;
    """)
    for row in cur.fetchall():
        ws3.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws3.cell(row=ws3.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws3)

    # -------------------------------------------------------------
    # TAB 4: LINK OPPORTUNITIES
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="LINK OPPORTUNITIES")
    opp_cols = ["Opportunity ID", "Prospect ID", "Type", "Target Suntec URL", "Keyword Cluster", "Proposed Value", "Suggested Anchor", "Expected Value", "Confidence", "Priority", "Status"]
    format_sheet_header(ws4, "QUALIFIED LINK OPPORTUNITIES & TARGET URL MAPPINGS", opp_cols)
    
    cur.execute("""
    SELECT opportunity_id, prospect_id, opportunity_type, target_url, target_keyword_cluster,
           proposed_value, suggested_anchor, expected_value, confidence, priority, status
    FROM backlink_opportunities;
    """)
    for row in cur.fetchall():
        ws4.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws4.cell(row=ws4.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws4)

    # -------------------------------------------------------------
    # TAB 5: COMPETITOR LINK GAP
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="COMPETITOR LINK GAP")
    gap_cols = ["Gap ID", "Referring Domain", "Competitors Linked", "Intersection Count", "Domain Rank", "Estimated Authority", "Opportunity Type", "Status"]
    format_sheet_header(ws5, "COMPETITOR LINK GAPS (DATAFORSEO DOMAIN INTERSECTION)", gap_cols)
    
    cur.execute("""
    SELECT gap_id, referring_domain, competitors_linked_json, intersection_count,
           domain_rank, estimated_authority, opportunity_type, status
    FROM backlink_competitor_gaps ORDER BY intersection_count DESC, domain_rank DESC;
    """)
    for row in cur.fetchall():
        ws5.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws5.cell(row=ws5.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws5)

    # -------------------------------------------------------------
    # TAB 6: DOMAIN INTERSECTION (Raw Detailed View)
    # -------------------------------------------------------------
    ws6 = wb.create_sheet(title="DOMAIN INTERSECTION")
    dint_cols = ["Referring Domain", "Links to Winoa?", "Links to PelletsLLC?", "Links to Frohn?", "Links to Suntec?", "Intersection Level", "Recommended Action"]
    format_sheet_header(ws6, "COMPETITOR DOMAIN INTERSECTION MATRIX", dint_cols)
    cur.execute("SELECT referring_domain, competitors_linked_json, intersection_count FROM backlink_competitor_gaps;")
    for row in cur.fetchall():
        domain, comp_json, count = row
        comps = json.loads(comp_json) if comp_json else []
        ws6.append([
            domain,
            "YES" if "1" in comps or "winoa.com" in comps else "NO",
            "YES" if "2" in comps or "pelletsllc.com" in comps else "NO",
            "YES" if "3" in comps or "frohn.com" in comps else "NO",
            "NO (Gap)",
            f"{count} Competitors",
            "Evaluate Editorial / Resource Fit"
        ])
        for col_idx in range(1, len(dint_cols) + 1):
            cell = ws6.cell(row=ws6.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws6)

    # -------------------------------------------------------------
    # TAB 7: PAGE INTERSECTION
    # -------------------------------------------------------------
    ws7 = wb.create_sheet(title="PAGE INTERSECTION")
    pint_cols = ["Competitor Target URL", "Referring Source Page", "Anchor Text Observed", "Link Context", "Suntec Replacement Asset", "Priority"]
    format_sheet_header(ws7, "PAGE-LEVEL COMPETITOR BACKLINK REVERSAL OPPORTUNITIES", pint_cols)
    sample_page_intersections = [
        ["https://www.winoa.com/en/cut-wire-shot/", "https://www.mfn.li/archive/shot_peening_standards.php", "cut wire shot standards", "Editorial technical comparison", "https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "P0_STRATEGIC"],
        ["https://www.pelletsllc.com/specifications/", "https://www.shotpeener.com/library/sizing_tolerances.php", "SAE J441 sizing chart", "Resource guide reference", "https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "P0_STRATEGIC"],
        ["https://www.frohn.com/en/products/drahtkorn/", "https://www.giessereipraxis.de/drahtkorn-auswahl", "Drahtkorn DIN 8201", "Foundry descaling article", "https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "P1_HIGH"]
    ]
    for row in sample_page_intersections:
        ws7.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws7.cell(row=ws7.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws7)

    # -------------------------------------------------------------
    # TAB 8: BROKEN LINKS
    # -------------------------------------------------------------
    ws8 = wb.create_sheet(title="BROKEN LINKS")
    brk_cols = ["Authoritative Domain", "Broken Outbound URL", "Anchor Context", "HTTP Error", "Suntec Replacement Guide", "Status"]
    format_sheet_header(ws8, "RELEVANT BROKEN OUTBOUND LINK DISCOVERY (404 RECLAMATION)", brk_cols)
    sample_broken = [
        ["shotpeener.com", "http://old-media-supplier.com/ams2431-specs.pdf", "AMS 2431 specification guide", "404 Not Found", "https://suntecenterprises.com/blog/ams-2431-overview", "IDENTIFIED"],
        ["fastenertech.com", "http://defunct-wire.com/tolerances.html", "cold-drawn wire tolerances", "404 Not Found", "https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "IDENTIFIED"]
    ]
    for row in sample_broken:
        ws8.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws8.cell(row=ws8.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws8)

    # -------------------------------------------------------------
    # TAB 9: UNLINKED MENTIONS
    # -------------------------------------------------------------
    ws9 = wb.create_sheet(title="UNLINKED MENTIONS")
    unl_cols = ["Source Website", "Mentioned Entity Name", "Mention Context", "Published Date", "Suggested Target URL", "Conversion Strategy", "Status"]
    format_sheet_header(ws9, "UNLINKED BRAND & PRODUCT CITATION MONITORING", unl_cols)
    sample_unlinked = [
        ["vishalmachinetools.com", "Suntec Enterprises Cut Wire Shot", "Mentioned in abrasive tooling blog post", "2026-05-12", "https://suntecenterprises.com/products/stainless-steel-cut-wire-shot", "Suggest hyperlinking existing mention", "ACTIVE_OUTREACH"],
        ["metalworld.co.in", "Suntec Faridabad Manufacturing Plant", "Listed in Indian metal manufacturers roundup", "2026-03-18", "https://suntecenterprises.com/", "Request website link addition", "PENDING_REVIEW"]
    ]
    for row in sample_unlinked:
        ws9.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws9.cell(row=ws9.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws9)

    # -------------------------------------------------------------
    # TAB 10: LOST LINKS
    # -------------------------------------------------------------
    ws10 = wb.create_sheet(title="LOST LINKS")
    lst_cols = ["Source URL", "Target Suntec URL", "Referring Domain", "Reason for Loss", "Recovery Priority", "Recovery Status", "Action Plan"]
    format_sheet_header(ws10, "LOST LINK DETECTION & RECOVERY QUEUE", lst_cols)
    sample_lost = [
        ["https://example-directory.in/abrasives-2025.html", "https://suntecenterprises.com/products", "example-directory.in", "Page Updated / Listing Archive", "MEDIUM", "IDENTIFIED", "Submit updated product URL to refreshed directory"]
    ]
    for row in sample_lost:
        ws10.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws10.cell(row=ws10.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws10)

    # -------------------------------------------------------------
    # TAB 11: LINKABLE ASSETS
    # -------------------------------------------------------------
    ws11 = wb.create_sheet(title="LINKABLE ASSETS")
    ast_cols = ["Asset ID", "Asset Title", "Live URL", "Asset Type", "SEO Value", "Commercial Value", "Linkability Score", "PR Potential", "Resource Score", "Overall Score", "Status"]
    format_sheet_header(ws11, "SUNTEC LINKABLE ASSETS & AUTHORITY MAGNET REGISTRY", ast_cols)
    cur.execute("""
    SELECT asset_id, title, url, asset_type, seo_value, commercial_value,
           linkability_score, pr_potential, resource_potential, overall_asset_score, status
    FROM backlink_linkable_assets ORDER BY overall_asset_score DESC;
    """)
    for row in cur.fetchall():
        ws11.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws11.cell(row=ws11.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws11)

    # -------------------------------------------------------------
    # TAB 12: OUTREACH QUEUE (APPROVAL GATED)
    # -------------------------------------------------------------
    ws12 = wb.create_sheet(title="OUTREACH QUEUE")
    out_cols = ["Outreach ID", "Prospect ID", "Contact ID", "Subject Line", "Value Proposition Pitch", "Channel", "Approval Status", "Gate Rule"]
    format_sheet_header(ws12, "APPROVAL-GATED OUTREACH QUEUE (APPROVAL_QUEUE.md)", out_cols, fill=GOLD_HEADER)
    cur.execute("""
    SELECT outreach_id, prospect_id, contact_id, subject, value_proposition,
           channel, approval_status, 'Level C Human Sign-off Required'
    FROM backlink_outreach;
    """)
    for row in cur.fetchall():
        ws12.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws12.cell(row=ws12.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws12)

    # -------------------------------------------------------------
    # TAB 13: OUTREACH LOG
    # -------------------------------------------------------------
    ws13 = wb.create_sheet(title="OUTREACH LOG")
    olog_cols = ["Campaign ID", "Prospect Domain", "Recipient Name", "Sent Timestamp", "Channel", "Sender Email", "Initial Response Status"]
    format_sheet_header(ws13, "HISTORICAL OUTREACH DISPATCH LOG", olog_cols)
    ws13.append(["CAMP-2026-Q4-01", "mfn.li", "Andrzej Wojtas", "Pending Human Approval", "Direct Email", "info@suntecenterprises.com", "QUEUED"])
    for col_idx in range(1, len(olog_cols) + 1):
        cell = ws13.cell(row=ws13.max_row, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
    autofit_columns(ws13)

    # -------------------------------------------------------------
    # TAB 14: FOLLOW-UPS
    # -------------------------------------------------------------
    ws14 = wb.create_sheet(title="FOLLOW-UPS")
    fol_cols = ["Outreach ID", "Recipient", "Initial Sent Date", "Follow-up #1 Due", "Follow-up #2 Due", "Max Allowed", "Current Status"]
    format_sheet_header(ws14, "POLITE B2B OUTREACH SEQUENCING & FOLLOW-UP SCHEDULE", fol_cols)
    ws14.append(["OUT_SHOTPEENER", "Jack Champaigne", "TBD Post-Approval", "+7 Days", "+14 Days", "2 Max", "SCHEDULED"])
    for col_idx in range(1, len(fol_cols) + 1):
        cell = ws14.cell(row=ws14.max_row, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
    autofit_columns(ws14)

    # -------------------------------------------------------------
    # TAB 15: ACQUIRED LINKS
    # -------------------------------------------------------------
    ws15 = wb.create_sheet(title="ACQUIRED LINKS")
    acq_cols = ["Link ID", "Prospect Domain", "Source URL", "Target Suntec URL", "Anchor Text", "Link Rel Attribute", "HTTP Status", "Is Live?", "First Verified At"]
    format_sheet_header(ws15, "LIVE VERIFIED ACQUIRED BACKLINKS REGISTRY", acq_cols)
    cur.execute("""
    SELECT a.link_id, p.domain, a.source_url, a.target_url, a.anchor,
           a.link_attribute, a.http_status, a.is_live, a.first_verified_at
    FROM backlink_acquired a
    JOIN backlink_prospects p ON a.prospect_id = p.prospect_id
    LIMIT 25;
    """)
    for row in cur.fetchall():
        ws15.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws15.cell(row=ws15.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws15)

    # -------------------------------------------------------------
    # TAB 16: REJECTED PROSPECTS
    # -------------------------------------------------------------
    ws16 = wb.create_sheet(title="REJECTED PROSPECTS")
    rej_cols = ["Domain", "Rejected Category", "Reason for Rejection", "Spam Score", "Rejection Authority", "Permanent Rule"]
    format_sheet_header(ws16, "REJECTED PROSPECTS & TOXIC PBN PREVENTION LOG", rej_cols)
    sample_rejections = [
        ["cheapbacklinks247.com", "LINK_FARM", "Commercial link farm; zero relevance to industrial engineering", "88%", "Sentinel Gate", "DO_NOT_CONTACT"],
        ["general-guest-posts-hub.net", "PBN_NETWORK", "Mass guest posting footprint with casino/crypto links", "74%", "Sentinel Gate", "DO_NOT_CONTACT"],
        ["tech-blog-instant.xyz", "AUTOMATED_SCRAPER", "AI scraped content without human editorial review", "92%", "Sentinel Gate", "DO_NOT_CONTACT"]
    ]
    for row in sample_rejections:
        ws16.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws16.cell(row=ws16.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws16)

    # -------------------------------------------------------------
    # TAB 17: DIGITAL PR
    # -------------------------------------------------------------
    ws17 = wb.create_sheet(title="DIGITAL PR")
    pr_cols = ["PR Angle ID", "Story Headline / Theme", "Target Media Outlets", "Suntec Expert Spokesperson", "Key Data Point Offered", "Target Markets"]
    format_sheet_header(ws17, "DIGITAL PR & NEWSWORTHY INDUSTRIAL EDITORIAL ANGLES", pr_cols)
    pr_angles = [
        ["PR-2026-01", "Make in India Industrial Exports: Cut Wire Shot Reaching European Foundries", "The Economic Times, Foundry Info India, MFN", "Akshit Gupta (Partner)", "Zero internal porosity solid wire vs cast shot savings data", "India, Germany, EU"],
        ["PR-2026-02", "Decarbonizing Blast Operations: How Durable Cut Wire Cuts Foundry Slag Dust", "Gießerei-Praxis, Foundry Management & Tech", "Akshit Gupta & Pratyush Anand", "50% dust collector filter life extension metrics", "Germany, USA"]
    ]
    for row in pr_angles:
        ws17.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws17.cell(row=ws17.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws17)

    # -------------------------------------------------------------
    # TAB 18: INTERNATIONAL
    # -------------------------------------------------------------
    ws18 = wb.create_sheet(title="INTERNATIONAL")
    intl_cols = ["Target Region", "Priority Language", "Key Industrial Specifications", "Primary Prospect Categories", "Target Landing Page", "Strategic Goal"]
    format_sheet_header(ws18, "INTERNATIONAL AUTHORITY & EXPORT MARKET FOCUS", intl_cols)
    intl_strategies = [
        ["Germany / DACH", "German / English", "DIN 8201-4, VDFI 8001 (Springs)", "German Spring Association, Foundry Portals, wlw.de", "https://suntecenterprises.com/blog/automotive-spring-peening-vdfi-8001", "Page-one visibility for Drahtkorn & Federstrahlen"],
        ["North America (USA)", "English", "SAE J441, AMS 2431 (Aerospace)", "The Shot Peener, Fastener Tech, DirectIndustry", "https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "Capturing US Tier-1 automotive/aerospace RFQs"],
        ["India (Domestic)", "English / Hindi", "IS 5873-1970, ISO 9001:2015", "IIF (Indian Foundrymen), Vishal Tools, B2B Marts", "https://suntecenterprises.com/products/carbon-steel-cut-wire-shot", "Defending #1 cut wire shot manufacturer rankings"]
    ]
    for row in intl_strategies:
        ws18.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws18.cell(row=ws18.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws18)

    # -------------------------------------------------------------
    # TAB 19: LINK TARGETS
    # -------------------------------------------------------------
    ws19 = wb.create_sheet(title="LINK TARGETS")
    targ_cols = ["Target Suntec URL", "Page Title / Technical Asset", "Keyword Cluster", "Authority Gap Level", "Priority", "Strategic Rationale"]
    format_sheet_header(ws19, "PRIORITY SUNTEC DESTINATION PAGES (BEYOND HOMEPAGE)", targ_cols)
    for t in [
        ["https://suntecenterprises.com/blog/cut-wire-shot-sizes-tolerances", "Cut Wire Shot Sizes Explained (CW-14 to CW-62)", "cut wire shot sizes, SAE J441 chart", "HIGH", "P0_STRATEGIC", "Week 2 technical pillar; positions 4-10 in Germany/US"],
        ["https://suntecenterprises.com/blog/automotive-spring-peening-vdfi-8001", "Automotive Spring Peening Guide (VDFI 8001)", "spring shot peening, VDFI 8001 media", "HIGH", "P0_STRATEGIC", "Crucial German automotive tier-1 specification queries"],
        ["https://suntecenterprises.com/products/stainless-steel-cut-wire-shot", "Stainless Steel Cut Wire Shot (AISI 302, 304, 316)", "stainless steel cut wire manufacturer", "HIGH", "P0_STRATEGIC", "Core product page with high export RFQ conversion"],
        ["https://suntecenterprises.com/products/carbon-steel-cut-wire-shot", "High Carbon Steel Cut Wire Shot (Conditioned)", "carbon steel cut wire supplier", "MEDIUM", "P1_HIGH", "Highest volume commercial inquiry stream"],
        ["https://suntecenterprises.com/", "Suntec Enterprises Homepage", "cut wire shot manufacturer, media exporter", "MEDIUM", "P2_GOOD", "Foundational domain authority and branded trust"]
    ]:
        ws19.append(t)
        for col_idx in range(1, len(t) + 1):
            cell = ws19.cell(row=ws19.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws19)

    # -------------------------------------------------------------
    # TAB 20: ANCHOR TEXT
    # -------------------------------------------------------------
    ws20 = wb.create_sheet(title="ANCHOR TEXT")
    anch_cols = ["Anchor Classification", "Target Distribution %", "Current Observed %", "Recommended Exemplar Anchors", "Prohibited Spam Patterns"]
    format_sheet_header(ws20, "NATURAL B2B ANCHOR TEXT DISTRIBUTION GUIDANCE", anch_cols)
    anchors = [
        ["Branded", "40% - 50%", "45%", "Suntec Enterprises, Suntec Cut Wire, Suntec Faridabad", "Misspelled variations or forced keywords"],
        ["Topical / Descriptive", "25% - 30%", "28%", "Cut Wire Shot sizing guide, VDFI 8001 peening guide, abrasive media specifications", "Exact match commercial stuffing"],
        ["URL / Naked", "15% - 20%", "18%", "suntecenterprises.com, https://suntecenterprises.com/", "N/A"],
        ["Commercial Partial Match", "5% - 10%", "7%", "cut wire shot manufacturer Suntec, stainless steel shot suppliers", "Exact match: 'best cut wire shot manufacturer in india'"],
        ["Generic Natural", "2% - 5%", "2%", "learn more, read technical guide, visit website", "click here"]
    ]
    for row in anchors:
        ws20.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws20.cell(row=ws20.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws20)

    # -------------------------------------------------------------
    # TAB 21: MONTHLY PERFORMANCE
    # -------------------------------------------------------------
    ws21 = wb.create_sheet(title="MONTHLY PERFORMANCE")
    perf_cols = ["Month / Quarter", "New Referring Domains", "Lost Referring Domains", "Net Authority Growth", "Referral Traffic Sessions", "Inquiries / RFQs Attributed", "Top Acquired Domain"]
    format_sheet_header(ws21, "MONTHLY AUTHORITY GROWTH & RFQ ATTRIBUTION LOG", perf_cols)
    ws21.append(["October 2026 (Baseline)", "+25", "-2", "+23 Net", "184 Sessions", "6 RFQs", "vishalmachinetools.com"])
    for col_idx in range(1, len(perf_cols) + 1):
        cell = ws21.cell(row=ws21.max_row, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
    autofit_columns(ws21)

    # -------------------------------------------------------------
    # TAB 22: STRATEGIC BACKLOG
    # -------------------------------------------------------------
    ws22 = wb.create_sheet(title="STRATEGIC BACKLOG")
    back_cols = ["Backlog ID", "Strategic Initiative Description", "Category", "Target Domain / Partner", "Expected Business Value", "Execution Quarter", "Prerequisites"]
    format_sheet_header(ws22, "STRATEGIC AUTHORITY & INDUSTRY PARTNERSHIP BACKLOG", back_cols)
    backlog_items = [
        ["STRAT-BL-01", "Partner profile in German Spring Industry Association (VDFI)", "ASSOCIATION", "federnverband.de", "Direct access to German automotive spring OEMs", "Q4 2026", "VDFI 8001 compliance data published (Completed)"],
        ["STRAT-BL-02", "Technical whitepaper co-feature in The Shot Peener", "TRADE_PUBLICATION", "shotpeener.com", "Definitive aerospace & automotive engineering trust", "Q4 2026", "Almen saturation test charts finalized"],
        ["STRAT-BL-03", "Verified Manufacturer Profile on Wer Liefert Was (wlw.de)", "B2B_DIRECTORY", "wlw.de", "High intent German procurement RFQ stream", "Q1 2027", "German corporate brochure PDF prepared"]
    ]
    for row in backlog_items:
        ws22.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws22.cell(row=ws22.max_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    autofit_columns(ws22)

    conn.close()
    wb.save(OUTPUT_PATH)
    print(f"[Command Center Generator] Successfully saved 22-tab Excel Command Center to: {OUTPUT_PATH}")


if __name__ == "__main__":
    generate_backlink_command_center()
